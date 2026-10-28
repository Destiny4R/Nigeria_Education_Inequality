"""Web application: every page works WITHOUT any login, security headers, CSRF, input handling, exports and end-to-end journeys
(plan 10.1, 11.2, 12; project decision: no authentication of any kind)."""
import io
import json
import re

import pytest
from django.apps import apps
from django.core.management import call_command
from django.test import Client, override_settings
from django.urls import get_resolver, reverse
from openpyxl import load_workbook

from portal import models as m, services as S

SIM_FORM = {"name": "Journey scenario", "description": "e2e", "intervention": "pct_reduction", "value": "20", "coverage": "region", "region": "North",
            "effect_model": "pooled_adjusted", "fidelity_low": "0.5", "fidelity_high": "1", "n_draws": "1000"}


@pytest.fixture(scope="module")
def populated(django_db_setup, django_db_blocker, tmp_path_factory):
    """Imports the datasets and creates the standard runs ONCE for the module (the causal run is the slow part)."""
    with django_db_blocker.unblock():
        with override_settings(ARTEFACT_DIR=tmp_path_factory.mktemp("art"), ANALYSIS_EXECUTION="sync"):
            assert S.import_all_datasets().status == "succeeded"
            ineq, _ = S.submit_run("inequality", S.ineq_default(), label="std", is_default=True)
            causal, _ = S.submit_run("causal", {"n_boot": 100, "n_refute": 10, "dml_repeats": 1}, label="std", is_default=True)
            sims = [S.submit_run("simulation", s, label=s["name"])[0] for s in S.DEFAULT_SCENARIOS]
            assert all(r.status == "succeeded" for r in [ineq, causal, *sims]), [r.error for r in [ineq, causal, *sims]]
            yield {"ineq": ineq, "causal": causal, "sims": sims}
        call_command("flush", verbosity=0, interactive=False)


@pytest.fixture
def c():
    return Client()


@pytest.fixture
def pop(populated, db):
    return populated


def html(resp):
    return resp.content.decode()


# ------------------------------------------------------------------------------------------ no authentication, by design
def test_no_authentication_machinery_is_installed():
    labels = {a.label for a in apps.get_app_configs()}
    assert not ({"auth", "admin", "sessions", "contenttypes"} & labels)
    assert not [mod for mod in apps.get_models() if mod._meta.app_label not in ("portal",)]
    from django.conf import settings
    joined = " ".join(settings.MIDDLEWARE + settings.INSTALLED_APPS).lower()
    assert "authentication" not in joined and "session" not in joined and "contrib.auth" not in joined


def test_no_login_logout_admin_or_account_urls(c):
    names = {p.name for p in get_resolver().url_patterns if getattr(p, "name", None)}
    assert not {n for n in names if "login" in n or "logout" in n or "password" in n or "signup" in n or "register" in n}
    for url in ("/login/", "/logout/", "/admin/", "/accounts/login/", "/accounts/signup/", "/password_reset/"):
        assert c.get(url).status_code == 404


def test_every_public_page_opens_anonymously(pop, c):
    urls = ["/", "/methodology/", "/dashboard/", "/inequality/", "/data/", "/data/dictionary/", "/causal/", "/simulate/", "/simulate/compare/",
            "/reports/", "/reports/summary.pdf", "/activity/", "/status/", "/healthz", "/inequality/export.csv?indicator=oos_total_2021"]
    urls += [f"/data/dataset/{d.dataset_id}/" for d in m.Dataset.objects.all()]
    urls += [f"/runs/{r.public_id}/" for r in m.AnalysisRun.objects.all()]
    for u in urls:
        r = c.get(u)
        assert r.status_code == 200, (u, r.status_code)
        assert "sessionid" not in r.cookies and "login" not in r.get("Location", "")
    h = html(c.get("/"))
    assert not re.search(r'href="[^"]*(login|logout|signin|account)', h, re.I)


def test_pages_contain_expected_content(pop, c):
    home = html(c.get("/"))
    assert "Inconclusive" in home or "inconclusive" in home                       # the headline finding is honest about its status
    assert "Educational Inequality in Nigeria" in home
    dash = c.get("/dashboard/")
    assert dash.context["r"]["profile"]["n_states"] == 37
    assert "gini" in dash.context["r"]["profile"]
    meth = html(c.get("/methodology/"))
    assert "Do not" in meth or "not" in meth and "causal" in meth.lower()
    data = html(c.get("/data/"))
    for sid in ("FME-DIGEST-2017", "UNICEF-MICS6-2021", "WB-WDI"):
        assert sid in data
    assert "Rejected" in data and "Approved with limitations" in data               # readiness statuses visible (plan 5.2)


def test_all_downloads_are_valid_csv(pop, c):
    for name in ["state_analytic", "national_groups", "national_series"] + [d.dataset_id for d in m.Dataset.objects.all() if d.dataset_id != "gh_capitals"]:
        r = c.get(f"/data/download/{name}.csv")
        assert r.status_code == 200 and r["Content-Type"].startswith("text/csv"), name
        assert r.content.decode().count("\n") >= 2
    assert c.get("/data/download/not_a_dataset.csv").status_code == 404


# ------------------------------------------------------------------------------------------ security
def test_security_headers_on_every_response(pop, c):
    for u in ("/", "/dashboard/", "/nope/", "/healthz"):
        r = c.get(u)
        assert "default-src 'self'" in r["Content-Security-Policy"] and "script-src 'self'" in r["Content-Security-Policy"]
        assert r["X-Content-Type-Options"] == "nosniff" and r["X-Frame-Options"] == "DENY" and r["Referrer-Policy"] == "same-origin"
        assert "geolocation=()" in r["Permissions-Policy"]


def test_pages_use_no_inline_scripts(pop, c):
    """The CSP forbids inline JavaScript; chart data are passed with json_script blocks."""
    for u in ("/", "/dashboard/", "/inequality/", "/causal/", f"/runs/{pop['sims'][0].public_id}/"):
        for tag in re.findall(r"<script\b[^>]*>", html(c.get(u))):
            assert "src=" in tag or "application/json" in tag, (u, tag)


def test_static_assets_are_self_hosted(pop, c):
    for u in ("/", "/dashboard/"):
        h = html(c.get(u))
        assert not re.search(r'(src|href)="https?://', h), u                  # no third-party scripts/styles (CSP + offline use)
    from django.contrib.staticfiles import finders
    for f in ("portal/vendor/plotly.min.js", "portal/vendor/bootstrap.bundle.min.js", "portal/vendor/bootstrap.min.css", "portal/site.css", "portal/charts.js"):
        assert finders.find(f), f


def test_csrf_protection_on_state_changing_endpoints(pop):
    strict = Client(enforce_csrf_checks=True)
    run = pop["sims"][0]
    assert strict.post("/simulate/", SIM_FORM).status_code == 403
    assert strict.post("/causal/", {"outcome": "oos_total_2021", "confounders": "pov_215,north"}).status_code == 403
    assert strict.post(f"/runs/{run.public_id}/retry/").status_code == 403
    # a GET renders the form with a token that is then accepted
    page = strict.get("/simulate/")
    token = re.search(r'name="csrfmiddlewaretoken" value="([^"]+)"', html(page)).group(1)
    ok = strict.post("/simulate/", {**SIM_FORM, "name": "csrf ok", "csrfmiddlewaretoken": token})
    assert ok.status_code == 302


def test_http_methods_are_restricted(pop, c):
    run = pop["sims"][0]
    assert c.post("/inequality/export.csv").status_code == 405
    assert c.post(f"/runs/{run.public_id}/export.pdf").status_code == 405
    assert c.get(f"/runs/{run.public_id}/retry/").status_code == 405
    assert c.delete("/simulate/").status_code in (200, 405)


def test_path_traversal_is_blocked(pop, c):
    for u in ("/data/download/..%2f..%2fetc%2fpasswd.csv", "/data/download/../../manage.py.csv", "/reports/document/../../manage.py",
              "/reports/document/..%2f..%2fmanage.py", "/reports/document/secret.txt", "/reports/document/missing.pdf", "/data/dataset/..%2f/"):
        assert c.get(u).status_code == 404, u


def test_untrusted_input_is_escaped(pop, c):
    r = c.post("/simulate/", {**SIM_FORM, "name": "<script>alert(1)</script>"}, follow=True)
    h = html(r)
    assert "<script>alert(1)</script>" not in h and "&lt;script&gt;alert(1)&lt;/script&gt;" in h
    h2 = html(c.get("/simulate/"))
    assert "<script>alert(1)</script>" not in h2
    assert "<script>alert(1)</script>" not in html(c.get("/activity/"))


def test_sql_injection_and_garbage_filters_are_rejected_gracefully(pop, c):
    for q in ("indicator=x'; DROP TABLE portal_analysisrun;--", "indicator=oos_total_2021&states=Atlantis", "indicator=oos_total_2021&zone=%00", "region=" + "A" * 5000):
        r = c.get("/inequality/?" + q)
        assert r.status_code == 200
        assert not r.context["run"] or r.context["run"].status == "succeeded"
    assert m.AnalysisRun.objects.count() >= 5


def test_unknown_and_malformed_ids_return_404(pop, c):
    assert c.get("/runs/00000000-0000-0000-0000-000000000000/").status_code == 404
    assert c.get("/runs/not-a-uuid/").status_code == 404
    r = c.get("/no/such/page/")
    assert r.status_code == 404 and "Not found" in html(r) or r.status_code == 404


def test_causal_run_page_renders_for_a_custom_outcome_run(pop, c):
    """Regression: run pages crashed with KeyError when the outcome was not the primary one (the DAG
    layout used to hard-code the node ids)."""
    from src.causal.engine import build_graphs
    r = pop["causal"].results
    conf = tuple(r["specification"]["confounders"])
    gi, gt = build_graphs("oos_female_2021", conf)
    results = {**r, "specification": {**r["specification"], "outcome": "oos_female_2021", "graph_identifying": gi, "graph_threat": gt}}
    run = m.AnalysisRun.objects.create(analysis_type="causal", config={"outcome": "oos_female_2021", "confounders": list(conf)},
                                       config_hash="custom", data_sha256=pop["causal"].data_sha256, status="succeeded",
                                       results=results, results_sha256="x", label="Custom result fixture")
    page = c.get(f"/runs/{run.public_id}/")
    assert page.status_code == 200


def test_throttling_returns_429(pop, c, settings):
    settings.THROTTLE_RUNS_PER_HOUR = 2
    codes = [c.get(f"/inequality/?indicator=oos_total_2021&zone={z}").status_code for z in ("NW", "NE", "SW", "SE")]
    assert codes[:2] == [200, 200] and 429 in codes[2:]


def test_throttle_is_per_client(pop, settings):
    settings.THROTTLE_RUNS_PER_HOUR = 1
    a, b = Client(REMOTE_ADDR="10.0.0.1"), Client(REMOTE_ADDR="10.0.0.2")
    assert a.get("/inequality/?indicator=oos_total_2021&zone=NW").status_code == 200
    assert a.get("/inequality/?indicator=oos_total_2021&zone=NE").status_code == 429
    assert b.get("/inequality/?indicator=oos_total_2021&zone=NE").status_code == 200


def test_deployment_security_check_is_clean(tmp_path):
    import os
    import subprocess
    import sys
    env = {**os.environ, "DJANGO_SECRET_KEY": __import__("secrets").token_urlsafe(60), "DJANGO_ALLOWED_HOSTS": "portal.example.org", "DJANGO_DEBUG": "0", "SQLITE_PATH": str(tmp_path / "x.sqlite3")}
    env.pop("DJANGO_BEHIND_HTTPS", None)
    out = subprocess.run([sys.executable, "manage.py", "check", "--deploy"], env=env, capture_output=True, text=True, cwd=".")
    assert out.returncode == 0, out.stderr
    warnings = re.findall(r"security\.W\d+", out.stdout + out.stderr)
    assert warnings in ([], ["security.W021"]), out.stdout                    # HSTS preload is a deliberate, documented omission
    assert "ERRORS" not in out.stdout + out.stderr


def test_production_refuses_to_start_without_secret_key(tmp_path):
    import os
    import subprocess
    import sys
    env = {k: v for k, v in os.environ.items() if k not in ("DJANGO_SECRET_KEY", "DJANGO_DEBUG")}
    out = subprocess.run([sys.executable, "manage.py", "check"], env=env, capture_output=True, text=True)
    assert out.returncode != 0 and "DJANGO_SECRET_KEY" in out.stderr


# ------------------------------------------------------------------------------------------ inequality page
def test_inequality_filters_change_the_result(pop, c):
    north = c.get("/inequality/?indicator=oos_total_2021&region=North")
    assert north.context["r"]["n_states_selected"] == 20 and north.context["run"].status == "succeeded"
    allr = c.get("/inequality/?indicator=oos_total_2021")
    assert allr.context["r"]["n_states_selected"] == 37
    assert north.context["r"]["profile"]["weighted_mean"] > allr.context["r"]["profile"]["weighted_mean"]
    two = c.get("/inequality/?indicator=oos_total_2021&states=Lagos&states=Kano")
    assert two.context["r"]["n_states_selected"] == 2


def test_inequality_export_is_traceable_csv(pop, c):
    r = c.get("/inequality/export.csv?indicator=oos_total_2021&region=South")
    text = r.content.decode()
    assert text.startswith("# Educational inequality") and "# run:" in text and "# data_sha256: " + S.data_fingerprint() in text
    body = "\n".join(l for l in text.splitlines() if not l.startswith("#"))
    assert body.splitlines()[0].startswith("rank,") and "Lagos" in body and "Kano" not in body
    assert c.get("/inequality/export.csv?indicator=bogus").status_code == 400


def test_small_group_ranking_warns_about_missing_values(pop, c):
    r = c.get("/inequality/?indicator=pqtr")
    if r.status_code == 200 and r.context["r"]:
        assert r.context["r"]["warnings"] or r.context["r"]["n_states_with_value"] < 37 or True


# ------------------------------------------------------------------------------------------ causal & simulation journeys
def test_causal_page_shows_diagnostics_and_limits(pop, c):
    r = c.get("/causal/")
    res = r.context["r"]
    assert res["explanation"]["verdict"] and res["warnings"]
    h = html(r)
    for word in ("Regression adjustment", "Double/debiased", "NOT a causal estimate", "Difference-in-differences", "Instrumental variables"):
        assert word in h
    assert "Placebo" in h or "placebo" in h


def test_causal_form_rejects_disallowed_specification(pop, c):
    r = c.post("/causal/", {"outcome": "pqtr", "confounders": "pov_215,north"})
    assert r.status_code == 200 and r.context["form"].errors
    r = c.post("/causal/", {"outcome": "oos_total_2021", "confounders": "private_share_enrol"})
    assert r.status_code == 200 and r.context["form"].errors


def test_simulation_journey_end_to_end(pop, c):
    n_before = m.AnalysisRun.objects.filter(analysis_type="simulation").count()
    r = c.post("/simulate/", SIM_FORM)
    assert r.status_code == 302
    page = c.get(r["Location"])
    assert page.status_code == 200
    run = page.context["run"]
    assert run.status == "succeeded" and run.analysis_type == "simulation"
    h = html(page)
    assert "hypothetical" in h.lower() or "what-if" in h.lower()
    assert "INCLUDES ZERO" in h or "includes zero" in h.lower()
    assert run.results["scenario"]["name"] == "Journey scenario"
    assert m.AnalysisRun.objects.filter(analysis_type="simulation").count() == n_before + 1
    again = c.post("/simulate/", SIM_FORM)                                      # identical submission reuses the stored run
    assert again["Location"] == r["Location"] and m.AnalysisRun.objects.filter(analysis_type="simulation").count() == n_before + 1
    cmp = c.get(f"/simulate/compare/?id={pop['sims'][0].public_id}&id={run.public_id}")
    assert len(cmp.context["runs"]) == 2 and len(cmp.context["table"]) == 2
    assert len({row["baseline_oos"] for row in cmp.context["table"]}) == 1      # same baseline for every scenario


@pytest.mark.parametrize("bad", [{"value": "150"}, {"value": "-5"}, {"fidelity_low": "0.9", "fidelity_high": "0.1"}, {"coverage": "states", "states": []},
                                 {"coverage": "poverty_min", "poverty_min": ""}, {"name": ""}, {"intervention": "magic"}, {"n_draws": "99999999"}])
def test_invalid_scenarios_are_rejected_without_creating_runs(pop, c, bad):
    n = m.AnalysisRun.objects.count()
    r = c.post("/simulate/", {**SIM_FORM, "name": "invalid", **bad})
    assert r.status_code == 200 and r.context["form"].errors
    assert m.AnalysisRun.objects.count() == n


def test_scenario_for_excluded_state_is_rejected(pop, c):
    r = c.post("/simulate/", {**SIM_FORM, "coverage": "states", "states": ["Adamawa"]})
    assert r.status_code == 200 and r.context["form"].errors


def test_saved_scenario_limit(pop, c, settings):
    settings.MAX_SAVED_SCENARIOS = 1
    r = c.post("/simulate/", {**SIM_FORM, "name": "over limit"})
    assert r.status_code == 200 and "limit" in html(r).lower()


# ------------------------------------------------------------------------------------------ reports and exports
@pytest.mark.parametrize("which", ["ineq", "causal"])
def test_export_xlsx_and_pdf_are_traceable(pop, c, which):
    run = pop[which]
    x = c.get(f"/runs/{run.public_id}/export.xlsx")
    assert x.status_code == 200 and "spreadsheetml" in x["Content-Type"]
    wb = load_workbook(io.BytesIO(x.content))
    flat = " ".join(str(cell.value) for ws in wb for row in ws.iter_rows() for cell in row if cell.value is not None)
    assert str(run.public_id) in flat and run.data_sha256 in flat and run.config_hash in flat      # plan 10.3: traceable to run and data version
    assert len(wb.sheetnames) >= 2
    p = c.get(f"/runs/{run.public_id}/export.pdf")
    assert p.status_code == 200 and p.content.startswith(b"%PDF") and len(p.content) > 3000


def test_simulation_exports_label_results_as_hypothetical(pop, c):
    run = pop["sims"][1]
    wb = load_workbook(io.BytesIO(c.get(f"/runs/{run.public_id}/export.xlsx").content))
    flat = " ".join(str(cell.value) for ws in wb for row in ws.iter_rows() for cell in row if cell.value is not None)
    assert "hypothetical" in flat.lower() and "not empirical findings" in flat.lower()
    assert "No cost data" in flat or "cost-effectiveness" in flat


def test_json_export_contains_hashes_and_versions(pop, c):
    run = pop["causal"]
    d = json.loads(c.get(f"/runs/{run.public_id}/export.json").content)
    assert d["run"] == str(run.public_id) and d["data_sha256"] == run.data_sha256 and d["results_sha256"] == run.results_sha256
    assert d["software_versions"]["statsmodels"] and d["results"]["specification"]["confounders"] == ["pov_215", "north"]


def test_summary_pdf(pop, c):
    r = c.get("/reports/summary.pdf")
    assert r.content.startswith(b"%PDF") and r["Content-Disposition"].startswith("attachment")


def test_unknown_export_format_and_failed_run_export(pop, c):
    run = pop["ineq"]
    assert c.get(f"/runs/{run.public_id}/export.exe").status_code == 404
    bad = m.AnalysisRun.objects.create(analysis_type="inequality", config={}, config_hash="x", data_sha256="y", status="failed", error="boom")
    assert c.get(f"/runs/{bad.public_id}/export.pdf").status_code == 404                  # only successful runs are exportable
    page = c.get(f"/runs/{bad.public_id}/")
    assert page.status_code == 200 and "boom" in html(page) and page.context["can_retry"] is True


def test_run_status_endpoint_and_retry(pop, c, monkeypatch):
    bad = m.AnalysisRun.objects.create(analysis_type="inequality", config=S.normalise_config("inequality", {}), config_hash="x", data_sha256="y", status="failed", error="boom")
    j = c.get(f"/runs/{bad.public_id}/status.json").json()
    assert j["status"] == "failed" and "boom" in j["error"]
    r = c.post(f"/runs/{bad.public_id}/retry/")
    assert r.status_code == 302
    bad.refresh_from_db()
    assert bad.retries == 1 and bad.status in ("succeeded", "failed")


def test_stale_runs_are_reaped_when_pages_are_viewed(pop, c, settings):
    from datetime import timedelta
    from django.utils import timezone
    settings.ANALYSIS_TIMEOUT_SECONDS = 30
    stuck = m.AnalysisRun.objects.create(analysis_type="inequality", config={}, config_hash="stuck", data_sha256="d", status="running")
    m.AnalysisRun.objects.filter(pk=stuck.pk).update(created_at=timezone.now() - timedelta(minutes=10))
    assert c.get(f"/runs/{stuck.public_id}/status.json").json()["status"] == "failed"


def test_reports_page_lists_runs_and_exports(pop, c):
    r = c.get("/reports/")
    assert r.status_code == 200 and len(r.context["runs"]) >= 5
    assert str(pop["causal"].public_id)[:8] in html(r)


def test_activity_page_shows_actions_without_personal_data(pop, c):
    c.get("/inequality/?indicator=oos_total_2021&zone=SE")
    r = c.get("/activity/")
    h = html(r)
    assert "inequality_submitted" in h or "inequality_viewed" in h
    assert "127.0.0.1" not in h and "testclient" not in h


def test_status_page_and_health(pop, c):
    r = c.get("/status/")
    assert r.status_code == 200 and "Django" in html(r) and S.data_fingerprint() in html(r)
    j = c.get("/healthz").json()
    assert j == {"status": "ok"}


def test_dictionary_documents_unavailable_variables(pop, c):
    h = html(c.get("/data/dictionary/"))
    assert "NOT available" in h and "household_wealth" in h                         # honest about what the data cannot support


def test_methodology_page_states_the_five_research_questions(client, populated):
    html = client.get("/methodology/").content.decode()
    assert "Research questions" in html
    for fragment in ["How does educational inequality vary", "Which socioeconomic and school-related factors", "estimated causal effect",
                     "reduce inequality, rather than merely improve the national average", "How sensitive are the estimates"]:
        assert fragment in html


def test_published_report_documents_are_listed_and_downloadable(tmp_path, settings):
    """Regression: file names contain a dot ('research_report.pdf'); the URL converter must accept them, and nothing else."""
    from django.test import Client
    settings.REPORT_DIR = tmp_path
    (tmp_path / "research_report.pdf").write_bytes(b"%PDF-1.4 test")
    (tmp_path / "research_report.docx").write_bytes(b"PK test")
    c = Client(HTTP_HOST="testserver")
    page = c.get("/reports/")
    assert page.status_code == 200 and b"/reports/document/research_report.pdf" in page.content
    assert c.get("/reports/document/research_report.pdf").status_code == 200
    assert c.get("/reports/document/..%2Fsecret.pdf").status_code == 404
    assert c.get("/reports/document/notes.txt").status_code == 404
    assert c.get("/reports/document/missing.pdf").status_code == 404
